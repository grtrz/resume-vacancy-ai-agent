from pathlib import Path

import fitz
import pytest
from docx import Document

from app.services.parsing import DocxParser, ParsingError, PdfParser, TextParser
from app.services.parsing.text_cleaner import clean_text

SAMPLE_LINES = [
    "Jane Doe",
    "Backend Engineer",
    "- Python",
    "- FastAPI",
    "Built APIs with PostgreSQL",
]


@pytest.fixture
def sample_resume_pdf(tmp_path: Path) -> Path:
    path = tmp_path / "sample_resume.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "\n".join(SAMPLE_LINES), fontsize=11)
    document.save(path)
    document.close()
    return path


@pytest.fixture
def sample_resume_docx(tmp_path: Path) -> Path:
    path = tmp_path / "sample_resume.docx"
    document = Document()
    document.add_paragraph("Jane Doe")
    document.add_paragraph("Backend Engineer")
    document.add_paragraph("Python", style="List Bullet")
    document.add_paragraph("FastAPI", style="List Bullet")
    document.add_paragraph("Built APIs with PostgreSQL")
    document.save(path)
    return path


@pytest.fixture
def sample_resume_txt(tmp_path: Path) -> Path:
    path = tmp_path / "sample_resume.txt"
    path.write_text("Jane Doe\r\n\r\n  -   Python\t\r\n- FastAPI\r\n", encoding="utf-8")
    return path


def test_pdf_parser_extracts_clean_text(sample_resume_pdf: Path) -> None:
    text = PdfParser().parse(sample_resume_pdf.read_bytes())

    assert "Jane Doe" in text
    assert "- Python" in text
    assert "Built APIs with PostgreSQL" in text


def test_docx_parser_extracts_paragraphs_and_bullets(sample_resume_docx: Path) -> None:
    text = DocxParser().parse(sample_resume_docx.read_bytes())

    assert text.splitlines() == SAMPLE_LINES


def test_text_parser_decodes_and_cleans_plain_text(sample_resume_txt: Path) -> None:
    text = TextParser().parse(sample_resume_txt.read_bytes())

    assert text == "Jane Doe\n\n- Python\n- FastAPI"


def test_clean_text_preserves_lines_and_normalizes_bullets() -> None:
    text = clean_text("  Jane\u00a0Doe\r\n\u2022   Python\r\n\r\n\r\nFastAPI  ")

    assert text == "Jane Doe\n- Python\n\nFastAPI"


def test_clean_text_removes_blank_lines_between_bullets() -> None:
    text = clean_text("Jane Doe\n\n- Python\n\n- FastAPI")

    assert text == "Jane Doe\n\n- Python\n- FastAPI"


def test_parsers_reject_empty_content() -> None:
    with pytest.raises(ParsingError):
        TextParser().parse(b"")
